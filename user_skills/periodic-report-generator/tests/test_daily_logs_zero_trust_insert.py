import contextlib
import importlib.util
import io
import sys
import tempfile
import types
import unittest
from pathlib import Path

import openpyxl


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "daily_logs_zero_trust_insert.py"


class DailyLogsZeroTrustInsertTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        sys.modules.setdefault("byted_aime_sdk", types.SimpleNamespace(call_aime_tool=lambda **kwargs: None))
        spec = importlib.util.spec_from_file_location("daily_logs_zero_trust_insert", SCRIPT_PATH)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        cls.module = module

    def make_xlsx(self, path: Path, sheet_name: str, headers: list[str], rows: list[list[object]]):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(path)

    def test_compute_stats_and_replace_placeholders(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx = Path(tmpdir) / "tasks.xlsx"
            self.make_xlsx(
                xlsx,
                "任务库",
                ["交付结果", "完成情况", "分类"],
                [
                    ["A", "进行中", "x"],
                    ["B", "准备中", "x"],
                    ["C", "已完成", "x"],
                    ["D", "暂停", "x"],
                    ["E", "", "x"],
                ],
            )
            stats = self.module.compute_task_status_stats(str(xlsx), "任务库")
            self.assertEqual(stats["opened"], 2)
            self.assertEqual(stats["completed"], 1)
            self.assertEqual(stats["paused"], 1)

            content = "今日开启：⚠️[数据断链_待自愈]\n今日完成：⚠️[数据断链_待自愈]\n今日暂停：⚠️[数据断链_待自愈]"
            resolved = self.module.replace_task_stats_placeholders(content, stats)
            self.assertIn("今日开启：2", resolved)
            self.assertIn("今日完成：1", resolved)
            self.assertIn("今日暂停：1", resolved)
            self.assertNotIn(self.module.UNRESOLVED_SENTINEL, resolved)

    def test_main_dry_run_resolves_content(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            task_xlsx = tmp / "task_stats.xlsx"
            daily_xlsx = tmp / "daily_logs.xlsx"
            self.make_xlsx(
                task_xlsx,
                "任务库",
                ["交付结果", "完成情况", "分类"],
                [["A", "进行中", "x"], ["B", "已完成", "x"]],
            )
            self.make_xlsx(
                daily_xlsx,
                "Daily_Logs",
                ["编号", "【日期】", "【日报内容】"],
                [["DL-20260528", "2026-05-28", "old"]],
            )

            original_download = self.module.mcp_download_lark_sheet
            original_existing_ids = self.module.list_existing_ids_via_cli
            original_argv = sys.argv[:]
            try:
                def fake_download(url: str):
                    if "wiki" in url:
                        return [str(task_xlsx)]
                    return [str(daily_xlsx)]

                self.module.mcp_download_lark_sheet = fake_download
                self.module.list_existing_ids_via_cli = lambda sheet_url, sheet_name: {"DL-20260528"}
                sys.argv = [
                    str(SCRIPT_PATH),
                    "--dry-run",
                    "--date",
                    "2026-05-29",
                    "--content",
                    "任务状态汇总：⚠️[数据断链_待自愈]\n正文",
                ]
                stdout = io.StringIO()
                with contextlib.redirect_stdout(stdout):
                    rc = self.module.main()
                output = stdout.getvalue()
                self.assertEqual(rc, 0)
                self.assertIn("开启 1 个，完成 1 个，暂停 0 个", output)
                self.assertIn("DL-20260529", output)
                self.assertNotIn(self.module.UNRESOLVED_SENTINEL, output)
            finally:
                self.module.mcp_download_lark_sheet = original_download
                self.module.list_existing_ids_via_cli = original_existing_ids
                sys.argv = original_argv
    def test_bracketed_headers_match_required_schema(self):
        headers = ["编号", "【日期】", "【日报内容】"]
        self.assertTrue(self.module.headers_match_required(headers))
        self.assertEqual(
            [self.module.normalize_header_name(h) for h in headers],
            ["编号", "日期", "日报内容"],
        )


if __name__ == "__main__":
    unittest.main()
